"""Resume text extraction (PDF / DOCX)."""

from __future__ import annotations

import io
from pathlib import Path

import fitz
from docx import Document

from app.domain.exceptions import EmptyResumeError, InvalidFileError
from app.utils.text import normalize_whitespace


class ResumeParser:
    def extract_text(self, file_bytes: bytes, filename: str | None = "resume.pdf") -> str:
        extension = Path(filename or "resume.pdf").suffix.lower()
        if extension == ".pdf":
            text = self._extract_pdf(file_bytes)
        elif extension == ".docx":
            text = self._extract_docx(file_bytes)
        else:
            raise InvalidFileError()

        if not text:
            raise EmptyResumeError()
        return text

    @staticmethod
    def _extract_pdf(pdf_bytes: bytes) -> str:
        try:
            document = fitz.open(stream=pdf_bytes, filetype="pdf")
        except Exception as exc:  # noqa: BLE001
            raise EmptyResumeError("Could not read the PDF resume.") from exc

        try:
            parts: list[str] = []
            for page in document:
                parts.append(page.get_text("text"))
            return normalize_whitespace("\n".join(parts))
        finally:
            document.close()

    @staticmethod
    def _extract_docx(docx_bytes: bytes) -> str:
        try:
            document = Document(io.BytesIO(docx_bytes))
        except Exception as exc:  # noqa: BLE001
            raise EmptyResumeError("Could not read the DOCX resume.") from exc

        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return normalize_whitespace("\n".join(parts))


def extract_text(file_bytes: bytes, filename: str | None = "resume.pdf") -> str:
    return ResumeParser().extract_text(file_bytes, filename)
